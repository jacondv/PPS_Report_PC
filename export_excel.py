from pathlib import Path
from openpyxl import load_workbook, Workbook


def merge_excel_folder(input_folder, output_file):
    input_folder = Path(input_folder)

    # tạo workbook tổng
    merged_wb = Workbook()
    merged_wb.remove(merged_wb.active)

    # duyệt toàn bộ file excel trong folder
    for file in input_folder.glob("*.xlsx"):

        # bỏ qua file output nếu nằm cùng folder
        if file.name == Path(output_file).name:
            continue

        wb = load_workbook(file)

        for sheet_name in wb.sheetnames:
            source_ws = wb[sheet_name]

            # tránh trùng tên sheet
            new_sheet_name = sheet_name
            counter = 1

            while new_sheet_name in merged_wb.sheetnames:
                new_sheet_name = f"{sheet_name}_{counter}"
                counter += 1

            target_ws = merged_wb.create_sheet(new_sheet_name)

            # copy dữ liệu
            for row in source_ws.iter_rows(values_only=True):
                target_ws.append(row)

    # lưu file
    merged_wb.save(output_file)

    print("Merge completed:", output_file)


if __name__ == "__main__":
    merge_excel_folder(r"C:\Users\danhv\Downloads\1input", r"C:\Users\danhv\Downloads\translated_qc_reports.xlsx")